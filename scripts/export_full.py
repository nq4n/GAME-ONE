import openpyxl
import json

wb = openpyxl.load_workbook('ECHOES_OF_SPACE_MASTER_GAME_BIBLE_2026.xlsx', data_only=True)

all_data = {}
for name in wb.sheetnames:
    sheet = wb[name]
    rows = []
    for r in range(1, sheet.max_row + 1):
        row_vals = [sheet.cell(r, c).value for c in range(1, sheet.max_column + 1)]
        if any(v is not None for v in row_vals):
            rows.append([str(v).strip() if v is not None else '' for v in row_vals])
    all_data[name] = rows

with open('all_sheets_full.json', 'w', encoding='utf-8') as f:
    json.dump(all_data, f, ensure_ascii=False, indent=2)

print(f"Dumped {len(all_data)} sheets to all_sheets_full.json")
